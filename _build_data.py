import json
import re
from pathlib import Path

import openpyxl

REGION_PREFIXES = [
    ("USA ", "USA"),
    ("England ", "England"),
    ("Team Canada", "Canada"),
    ("Norway ", "Norway"),
    ("Germany ", "Germany"),
    ("France ", "France"),
    ("Australia ", "Australia"),
    ("Finland ", "Finland"),
    ("Team Poland", "Poland"),
    ("Sweden ", "Sweden"),
    ("Spain ", "Spain"),
    ("Italy ", "Italy"),
    ("Belgian ", "Belgium"),
    ("South Africa ", "South Africa"),
    ("Austria ", "Austria"),
    ("The Netherlands ", "Netherlands"),
    ("Denmark ", "Denmark"),
    ("Portugal ", "Portugal"),
    ("Cymru ", "Wales"),
]

ATTACHMENT_RE = re.compile(
    r"^(CENTRAL HEAD|SMALLER HEAD|ARCANE HEADS|RANGE HEADS|MELEE HEADS|"
    r"ARCANE TAILS|RANGE TAILS|MELEE TAILS|RIGHT FOREARM|LEFT FOREARM|"
    r"RIGHT FIST|LEFT FIST|RIGHT ARM|LEFT ARM|DORSAL FEATURE|"
    r"HEAD|BACK|TAIL|WINGS|ARMS|CONDITIONING|DEFENSE OPTION|HEAVY WEAPON|"
    r"HEAVY WEAPON CRATE)\b",
    re.I,
)
TOTAL_RE = re.compile(r"TOTAL POINTS", re.I)
COMMAND_RE = re.compile(r"COMMAND CARD", re.I)
PC_CARD_RE = re.compile(r"^PC\s+CARD$", re.I)
FORMAT_RE = re.compile(r"^(Grand Melee|\d+\s*pts)\b", re.I)
CRATE_RE = re.compile(
    r"(Ammo Crate|Medical Crate|Fuel Canister|Mantlet|Heavy Weapon Crate|"
    r"^Blocker(\s+\d+)?$|^Skirmisher(\s+\d+)?$|^Raider(\s+\d+)?$|"
    r"^Weather Station|^Defenses$|^Trapdoor(\s+\d+)?$|"
    r"^HEAVY WEAPON\b|^WEAPON SUPPORT)",
    re.I,
)
# Army-wide ability cards that look like models in exports (Fane of Nyrro, etc.).
ARMY_ABILITY_RE = re.compile(r"^(Hunger)$", re.I)
COMPANION_RE = re.compile(
    r"^(Benkei|Sasha|Gallant|Invictus|Aberration|Wight|Drang|"
    r"War Boar(\s+MMD47)?|Ramhead|Corpse Crawler|"
    r"Mr\.?\s*Clogg|Phantasm|"
    r"Exponent Servitors|Permutation Servitors)$",
    re.I,
)
JUNK_PREFIX_RE = re.compile(
    r"^(Missing Card|Light Airdrop|FIN\s*-|BOOZE\s*-|FIGUREHEAD\s*-|"
    r"TENTACLE WEAPON|GUN EMPLACEMENTS|WEAPON SUPPORT)\b",
    re.I,
)


def is_skip_model(name: str) -> bool:
    return bool(
        CRATE_RE.search(name)
        or ARMY_ABILITY_RE.match(name)
        or COMPANION_RE.match(name)
        or JUNK_PREFIX_RE.match(name)
        or ATTACHMENT_RE.search(name)
    )


def is_free(points) -> bool:
    return points is None or points == 0


def looks_like_joke_title(name: str) -> bool:
    if len(name) > 48:
        return True
    if name.count(" ") >= 8:
        return True
    if re.search(r"[\"“”!?]{2,}|\.{3}|http|www\.", name, re.I):
        return True
    return False


def pick_caster(titles: list[str], entries: list[dict]) -> str:
    """Caster is the last free (no-cost) model before the first paid model.

    Warmachine exports start with junk titles / faction / PC CARD / crates, then
    the warcaster with no points, then costed models. Solos later in the list
    have a cost and must not be picked.
    """
    prefix: list[str] = []
    for title in titles:
        if not is_skip_model(title):
            prefix.append(title)
    for entry in entries:
        name = entry["name"]
        if entry.get("attachment") or is_skip_model(name):
            if not is_free(entry.get("points")) and not entry.get("attachment"):
                break
            continue
        if is_free(entry.get("points")):
            prefix.append(name)
            continue
        break

    candidates = [n for n in prefix if not looks_like_joke_title(n)]
    if candidates:
        return candidates[-1]
    if prefix:
        return prefix[-1]
    return next(
        (e["name"] for e in entries if not e.get("attachment") and not is_skip_model(e["name"])),
        titles[0] if titles else "Unknown",
    )


def region_for(team: str) -> str:
    for prefix, region in REGION_PREFIXES:
        if team.startswith(prefix):
            return region
    specials = {
        "Team Suokellopöllö": "Finland",
        "Brave Tin Soldiers": "Other",
        "Buccaneers Amatriciana": "Other",
        "Flemish Giant": "Belgium",
    }
    return specials.get(team, "Other")


def split_faction(faction: str) -> tuple[str, str]:
    if " - " in faction:
        army, theme = faction.split(" - ", 1)
        return army.strip(), theme.strip()
    return faction.strip(), ""


def collapse(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\xa0", " ")).strip()


def split_blocks(raw_lines: list[str | None]) -> list[list[str]]:
    blocks: list[list[str]] = []
    current: list[str] = []
    empty_run = 0
    for line in raw_lines:
        if line is None or not str(line).strip():
            empty_run += 1
            if current and empty_run >= 2 and any(TOTAL_RE.search(x) for x in current):
                blocks.append(current)
                current = []
            continue
        empty_run = 0
        current.append(str(line))
        if TOTAL_RE.search(str(line)):
            blocks.append(current)
            current = []
    if current:
        blocks.append(current)
    return [b for b in blocks if sum(1 for x in b if collapse(x)) >= 5]


def parse_block(lines: list[str], faction: str) -> dict | None:
    titles: list[str] = []
    entries: list[dict] = []
    commands: list[str] = []
    seen_pc = False
    seen_model = False
    in_commands = False
    faction_l = collapse(faction).lower() if faction else ""

    for raw in lines:
        text = collapse(raw)
        if not text or TOTAL_RE.search(text):
            continue
        if COMMAND_RE.search(text):
            in_commands = True
            continue
        if in_commands:
            commands.append(text)
            continue
        if PC_CARD_RE.match(text) or FORMAT_RE.match(text):
            if PC_CARD_RE.match(text):
                seen_pc = True
            continue
        if faction_l and text.lower() in {faction_l, faction_l.split(" - ")[0]}:
            continue

        points = None
        name = text
        match = re.match(r"^(\d+)\s+(.+)$", text)
        if match:
            points = int(match.group(1))
            name = match.group(2).strip()
            # Years / joke titles that happen to start with a number.
            if points >= 100 or (not seen_pc and not seen_model and looks_like_joke_title(name)):
                if not titles or titles[-1] != name:
                    titles.append(name)
                continue

        attachment = bool(ATTACHMENT_RE.search(name))
        if ARMY_ABILITY_RE.match(name):
            continue
        if points is not None and not attachment:
            seen_model = True
        if not seen_pc and not seen_model and points is None and not attachment:
            if not titles or titles[-1] != name:
                titles.append(name)
            continue
        entries.append({"name": name, "points": points, "attachment": attachment})

    if not entries and not titles:
        return None

    titled = [t for t in titles if not ARMY_ABILITY_RE.match(t) and not COMPANION_RE.match(t)]
    list_name = titled[0] if titled else (entries[0]["name"] if entries else "List")
    caster = pick_caster(titled, entries)

    return {
        "name": list_name,
        "caster": caster,
        "entries": entries,
        "commands": commands,
    }


def parse_player_lists(ws, col: int, faction: str) -> list[dict]:
    raw = [ws.cell(row, col).value for row in range(4, ws.max_row + 1)]
    lists = []
    for block in split_blocks(raw):
        parsed = parse_block(block, faction)
        if parsed:
            lists.append(parsed)
    return lists[:2]


def build_teams_json() -> None:
    wb = openpyxl.load_workbook(
        r"c:\Users\Rafał\Downloads\WTC Data by team.xlsx", data_only=True
    )

    teams = []
    two_lists = 0
    players_n = 0
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        players = []
        team_name = None
        for col in range(2, 7):
            pname = ws.cell(1, col).value
            team = ws.cell(2, col).value
            faction = ws.cell(3, col).value
            team_name = str(team).strip() if team else sheet_name
            faction_s = str(faction).strip() if faction else ""
            army, theme = split_faction(faction_s)
            lists = parse_player_lists(ws, col, faction_s)
            players_n += 1
            if len(lists) >= 2:
                two_lists += 1
            players.append(
                {
                    "id": f"{team_name}||{str(pname).strip() if pname else col}",
                    "name": str(pname).strip() if pname else f"Player {col-1}",
                    "faction": faction_s,
                    "army": army,
                    "theme": theme,
                    "lists": lists,
                }
            )
        teams.append(
            {
                "id": team_name,
                "name": team_name,
                "region": region_for(team_name),
                "players": players,
            }
        )

    out = Path(__file__).parent / "data" / "teams.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "event": "WTC",
        "source": "WTC Data by team.xlsx",
        "teams": teams,
    }
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {len(teams)} teams, {two_lists}/{players_n} players with 2 lists -> {out}")


if __name__ == "__main__":
    build_teams_json()
